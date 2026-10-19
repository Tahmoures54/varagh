# -*- coding: utf-8 -*-
"""سامانه کنترل حمل ورق فولادی فراسکو."""
from __future__ import annotations

import csv
import io
import json
import math
import os
import re
import secrets
import uuid
from datetime import date, datetime, time, timezone
from functools import wraps
from pathlib import Path
from urllib.parse import urlsplit

from flask import (
    Blueprint, Flask, abort, flash, make_response, redirect, render_template, request,
    send_from_directory, session, url_for,
)
from flask_login import LoginManager, current_user, login_required, login_user, logout_user
from sqlalchemy import inspect as sqlalchemy_inspect, or_, text
from sqlalchemy.exc import IntegrityError, OperationalError
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

from models import (
    AuditLog, Havaleh, HavalehItem, Inventory, SerialCounter, TransferControl, User, db,
    ROLE_ADMIN, ROLE_ANBAR, ROLE_HERASAT, ROLE_PROJECT, ROLE_VIEWER,
)

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

# Vercel function bundles are read-only; /tmp is writable but not persistent.
_default_data_dir = (
    Path("/tmp/varagh-data")
    if (os.environ.get("VERCEL") == "1" or os.environ.get("AWS_LAMBDA_FUNCTION_NAME"))
    else BASE_DIR / "data"
)
DATA_DIR = Path(os.environ.get("VARAGH_DATA_DIR") or _default_data_dir).resolve()
UPLOAD_DIR = Path(os.environ.get("VARAGH_UPLOAD_DIR") or (DATA_DIR / "receipts")).resolve()


def _normalize_database_url(url):
    """Normalize provider URLs for SQLAlchemy and the installed psycopg v3 driver."""
    if not url:
        return url
    url = url.replace("postgres://", "postgresql://", 1)
    if url.startswith("postgresql://"):
        url = url.replace("postgresql://", "postgresql+psycopg://", 1)
    return url


def _is_production_env():
    return (
        os.environ.get("FLASK_ENV") == "production"
        or os.environ.get("VERCEL") == "1"
        or bool(os.environ.get("AWS_LAMBDA_FUNCTION_NAME"))
    )


def _safe_next_url(candidate):
    """Return candidate only if it is a same-origin relative path."""
    if not candidate or not isinstance(candidate, str):
        return None
    if not candidate.startswith("/"):
        return None
    if candidate.startswith("//") or candidate.startswith("/\\"):
        return None
    parts = urlsplit(candidate)
    if parts.scheme or parts.netloc:
        return None
    return candidate


MAX_UPLOAD_BYTES = 8 * 1024 * 1024
ALLOWED_EXTENSIONS = {"pdf", "png", "jpg", "jpeg", "webp"}
STATUSES = ["ثبت‌شده", "در حال بررسی", "خارج‌شده", "ناقص", "نیازمند بررسی", "تکمیل‌شده"]
MOGHAYERAT = ["بدون مغایرت", "مغایرت مثبت", "مغایرت منفی", "نیازمند بررسی"]
BALE_KHEYR = ["خیر", "بله"]

# فیلدهایی که ممکن است در فرم حواله ارسال شوند
ALL_HAVALEH_FIELDS = {
    "tarikh_havaleh", "pelak", "noe_vasileh", "name_ranande", "mobile_ranande",
    "sharh_kala", "moshakhasat_varagh", "code_kala", "mogheiyat_varagh",
    "tedad_baste", "vazn_esnadi", "tarikh_vorud", "saat_vorud",
    "tarikh_khoruj", "saat_khoruj", "masoul_sabt", "vazn_por", "vazn_khali",
    "ghabz_por_file", "ghabz_khali_file", "vaziyat_havaleh", "vaziyat_moghayerat",
    "sharh_moghayerat", "masoul_peygiri", "tarikh_peygiri", "natije_peygiri",
    "taeed_nahayi", "tarikh_taeed", "tozihat",
}


def role_editable_fields(role):
    """فیلدهایی که این نقش مجاز به ویرایش است."""
    if role == ROLE_HERASAT:
        forbidden = {
            "vazn_esnadi", "vazn_por", "vazn_khali", "code_kala",
            "tedad_baste", "taeed_nahayi", "tarikh_taeed",
        }
    elif role == ROLE_ANBAR:
        forbidden = {
            "tarikh_vorud", "saat_vorud", "tarikh_khoruj", "saat_khoruj",
            "taeed_nahayi", "tarikh_taeed",
        }
    elif role == ROLE_PROJECT:
        forbidden = {
            "tarikh_vorud", "saat_vorud", "tarikh_khoruj", "saat_khoruj",
            "vazn_esnadi", "vazn_por", "vazn_khali", "code_kala", "tedad_baste",
        }
    else:
        forbidden = set()
    return ALL_HAVALEH_FIELDS - forbidden


def create_app(test_config=None):
    app = Flask(__name__)

    # ---------- SECRET_KEY (production must be explicit) ----------
    secret = os.environ.get("VARAGH_SECRET_KEY")
    if not secret:
        if _is_production_env():
            raise RuntimeError(
                "VARAGH_SECRET_KEY is required in production. "
                "Set a random 64-character secret before starting the app."
            )
        secret = secrets.token_hex(32)
        app.logger.warning(
            "VARAGH_SECRET_KEY is not set; using a temporary key. "
            "Sessions and CSRF tokens will be invalidated on restart."
        )

    app.config.update(
        SECRET_KEY=secret,
        SQLALCHEMY_DATABASE_URI=_normalize_database_url(os.environ.get(
            "DATABASE_URL", f"sqlite:///{(DATA_DIR / 'varagh.db').as_posix()}"
        )),
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        MAX_CONTENT_LENGTH=MAX_UPLOAD_BYTES,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "0") == "1",
        REMEMBER_COOKIE_HTTPONLY=True,
        REMEMBER_COOKIE_SAMESITE="Lax",
    )

    if test_config:
        app.config.update(test_config)

    # Engine pool options apply to PostgreSQL only; do not override tests.
    database_uri = app.config.get("SQLALCHEMY_DATABASE_URI", "")
    if "SQLALCHEMY_ENGINE_OPTIONS" not in app.config and database_uri.startswith(
        ("postgresql://", "postgresql+psycopg://")
    ):
        app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {
            "pool_pre_ping": True,
            "pool_recycle": 240,
            "pool_size": 1,
            "max_overflow": 0,
            "pool_timeout": 10,
            "connect_args": {"connect_timeout": 10},
        }

    try:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    except OSError:
        app.logger.warning("Data/upload directory is not writable: %s", DATA_DIR)

    db.init_app(app)

    login_manager = LoginManager()
    login_manager.login_view = "auth.login"
    login_manager.login_message = "برای ادامه ابتدا وارد سامانه شوید."
    login_manager.login_message_category = "warning"
    login_manager.init_app(app)

    @login_manager.user_loader
    def load_user(user_id):
        try:
            return db.session.get(User, int(user_id))
        except (TypeError, ValueError):
            return None

    # ---------- blueprints ----------
    auth_bp = Blueprint("auth", __name__)
    dashboard_bp = Blueprint("dashboard", __name__)
    havaleh_bp = Blueprint("havaleh", __name__)
    users_bp = Blueprint("users", __name__)
    inventory_bp = Blueprint("inventory", __name__)
    reports_bp = Blueprint("reports", __name__)

    # ---------- shared helpers ----------
    def log_action(action, entity=None, entity_id=None, detail=None):
        db.session.add(AuditLog(
            user_id=current_user.id if current_user.is_authenticated else None,
            action=action, entity=entity, entity_id=entity_id,
            detail=(detail or "")[:2000],
        ))

    def require_roles(*roles):
        def decorator(view):
            @wraps(view)
            @login_required
            def wrapped(*args, **kwargs):
                if current_user.role not in roles:
                    abort(403)
                return view(*args, **kwargs)
            return wrapped
        return decorator

    @app.before_request
    def csrf_protect():
        if request.method in ("POST", "PUT", "DELETE", "PATCH"):
            supplied = request.form.get("_csrf_token", "")
            expected = session.get("_csrf_token", "")
            if not expected or not supplied or not secrets.compare_digest(supplied, expected):
                abort(400, description="درخواست معتبر نیست؛ صفحه را تازه‌سازی و دوباره تلاش کنید.")

    @app.context_processor
    def inject_globals():
        if "_csrf_token" not in session:
            session["_csrf_token"] = secrets.token_urlsafe(32)
        return {"csrf_token": session["_csrf_token"], "current_user": current_user}

    @app.template_filter("fa_date")
    def fa_date(value):
        if not value:
            return "—"
        if isinstance(value, datetime):
            value = value.date()
        if not isinstance(value, date):
            return str(value)
        jy, jm, jd = gregorian_to_jalali(value.year, value.month, value.day)
        return f"{jy:04d}/{jm:02d}/{jd:02d}".translate(
            str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")
        )

    # ---------- form parsers ----------
    def parse_date(field):
        raw = (request.form.get(field) or "").strip()
        if not raw:
            return None
        try:
            return parse_flexible_date(raw)
        except ValueError as exc:
            raise ValueError(f"مقدار تاریخ «{field}» معتبر نیست: {exc}")

    def parse_time(field):
        raw = (request.form.get(field) or "").strip()
        if not raw:
            return None
        try:
            return time.fromisoformat(raw)
        except ValueError:
            raise ValueError(f"مقدار ساعت «{field}» معتبر نیست.")

    def parse_float(field, label, required=False):
        raw = (request.form.get(field) or "").strip()
        if not raw:
            if required:
                raise ValueError(f"فیلد «{label}» الزامی است.")
            return None
        try:
            value = float(raw.replace(",", ""))
        except ValueError:
            raise ValueError(f"مقدار «{label}» باید عددی باشد.")
        if not math.isfinite(value):
            raise ValueError(f"مقدار «{label}» نامعتبر است.")
        if value < 0:
            raise ValueError(f"مقدار «{label}» نمی‌تواند منفی باشد.")
        return value

    def parse_int(field, label, max_value=10_000_000):
        raw = (request.form.get(field) or "").strip()
        if not raw:
            return None
        try:
            value = int(raw)
        except ValueError:
            raise ValueError(f"مقدار «{label}» باید عدد صحیح باشد.")
        if value < 0:
            raise ValueError(f"مقدار «{label}» نمی‌تواند منفی باشد.")
        if value > max_value:
            raise ValueError(f"مقدار «{label}» بیش از حد مجاز است.")
        return value

    # ---------- voucher / receipt helpers ----------
    def next_voucher_number():
        jy, _, _ = gregorian_to_jalali(
            date.today().year, date.today().month, date.today().day
        )
        prefix = f"HV-{jy}-"
        counter = (
            SerialCounter.query.filter_by(prefix=prefix)
            .with_for_update()
            .first()
        )
        if counter is None:
            counter = SerialCounter(prefix=prefix, last_number=0)
            db.session.add(counter)
            db.session.flush()
        counter.last_number += 1
        return f"{prefix}{counter.last_number:04d}"

    def save_receipt(field, current_path=None):
        upload = request.files.get(field)
        if not upload or not upload.filename:
            return current_path

        # Extract extension from original name (may contain Persian characters)
        original = upload.filename
        if "." not in original:
            raise ValueError("فایل قبض باید پسوند داشته باشد.")
        extension = original.rsplit(".", 1)[-1].lower()
        if extension not in ALLOWED_EXTENSIONS:
            raise ValueError(
                "نوع فایل قبض مجاز نیست؛ فقط PDF، PNG، JPG یا WEBP پذیرفته می‌شود."
            )

        try:
            upload.stream.seek(0, 2)
            size = upload.stream.tell()
            upload.stream.seek(0)
        except (OSError, AttributeError):
            size = 0
        if size and size > MAX_UPLOAD_BYTES:
            raise ValueError("حجم فایل قبض نباید بیش از ۸ مگابایت باشد.")

        stored_name = f"{uuid.uuid4().hex}.{extension}"
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        upload.save(UPLOAD_DIR / stored_name)
        return stored_name

    # ---------- role / edit enforcement ----------
    def form_value_matches(field, current):
        raw = (request.form.get(field) or "").strip()
        if current is None:
            return raw == ""
        if isinstance(current, datetime):
            current = current.date()
        if isinstance(current, date):
            try:
                return parse_flexible_date(raw) == current
            except ValueError:
                return False
        if isinstance(current, time):
            return raw == current.strftime("%H:%M")
        if isinstance(current, bool):
            return (
                raw.lower() in ("1", "true", "on")
                if current
                else raw.lower() not in ("1", "true", "on")
            )
        if isinstance(current, (int, float)):
            if not raw:
                return False
            try:
                return float(raw) == float(current)
            except ValueError:
                return False
        return raw == str(current)

    def enforce_role_field_permissions(item, editable_fields):
        """فیلدهای محافظت‌شده که کاربر نمی‌تواند تغییر دهد نباید تغییر کنند."""
        guard_fields = {
            "tarikh_vorud": item.tarikh_vorud,
            "saat_vorud": item.saat_vorud,
            "tarikh_khoruj": item.tarikh_khoruj,
            "saat_khoruj": item.saat_khoruj,
            "vazn_esnadi": item.vazn_esnadi,
            "vazn_por": item.vazn_por,
            "vazn_khali": item.vazn_khali,
            "code_kala": item.code_kala,
            "tedad_baste": item.tedad_baste,
            "taeed_nahayi": item.taeed_nahayi,
            "tarikh_taeed": item.tarikh_taeed,
        }
        for field, current in guard_fields.items():
            if field in editable_fields:
                continue
            if field not in request.form:
                continue
            if not form_value_matches(field, current):
                raise ValueError(
                    f"نقش کاربری شما اجازهٔ تغییر فیلد «{field}» را ندارد."
                )

    # ---------- apply form to Havaleh ----------
    def apply_form(item, creating=False, editable_fields=None):
        if editable_fields is None:
            editable_fields = ALL_HAVALEH_FIELDS

        def can_edit(field):
            return field in editable_fields

        old_status = item.vaziyat_havaleh
        old_code = item.code_kala
        old_weight_esnadi = item.vazn_esnadi
        old_vazn_por = item.vazn_por
        old_vazn_khali = item.vazn_khali
        old_location = item.mogheiyat_varagh
        old_spec = item.moshakhasat_varagh

        # ----- lock after inventory deduction -----
        if item.inventory_deducted:
            requested_status = (
                request.form.get("vaziyat_havaleh") or old_status or ""
            ).strip()
            requested_code = (
                request.form.get("code_kala") or old_code or ""
            ).strip()

            requested_weight_esnadi = (
                parse_float("vazn_esnadi", "وزن اسنادی")
                if can_edit("vazn_esnadi") else old_weight_esnadi
            )
            requested_vazn_por = (
                parse_float("vazn_por", "وزن پُر باسکول")
                if can_edit("vazn_por") else old_vazn_por
            )
            requested_vazn_khali = (
                parse_float("vazn_khali", "وزن خالی باسکول")
                if can_edit("vazn_khali") else old_vazn_khali
            )
            requested_location = (
                ((request.form.get("mogheiyat_varagh") or "").strip()[:200] or None)
                if can_edit("mogheiyat_varagh") else old_location
            )
            requested_spec = (
                ((request.form.get("moshakhasat_varagh") or "").strip()[:200] or None)
                if can_edit("moshakhasat_varagh") else old_spec
            )

            diffs = (
                requested_status != (old_status or ""),
                requested_code != (old_code or ""),
                requested_weight_esnadi != old_weight_esnadi,
                requested_vazn_por != old_vazn_por,
                requested_vazn_khali != old_vazn_khali,
                requested_location != old_location,
                requested_spec != old_spec,
            )
            if any(diffs):
                raise ValueError(
                    "حواله‌ای که از موجودی آن کسر شده، قابل تغییر در وضعیت، "
                    "کد کالا، وزن‌ها، مشخصات یا محل موجودی نیست. "
                    "برای اصلاح، ابتدا با مدیر سامانه هماهنگ کنید."
                )

        # ----- write editable fields -----
        if can_edit("tarikh_havaleh"):
            item.tarikh_havaleh = parse_date("tarikh_havaleh")
        if can_edit("pelak"):
            item.pelak = (request.form.get("pelak") or "").strip()[:20] or None
        if can_edit("noe_vasileh"):
            item.noe_vasileh = (request.form.get("noe_vasileh") or "").strip()[:30] or None
        if can_edit("name_ranande"):
            item.name_ranande = (request.form.get("name_ranande") or "").strip()[:100] or None
        if can_edit("mobile_ranande"):
            item.mobile_ranande = (request.form.get("mobile_ranande") or "").strip()[:20] or None
        if can_edit("sharh_kala"):
            item.sharh_kala = (request.form.get("sharh_kala") or "").strip()[:100] or None
        if can_edit("moshakhasat_varagh"):
            item.moshakhasat_varagh = (request.form.get("moshakhasat_varagh") or "").strip()[:200] or None
        if can_edit("code_kala"):
            item.code_kala = (request.form.get("code_kala") or "").strip()[:30] or None
        if can_edit("mogheiyat_varagh"):
            item.mogheiyat_varagh = (request.form.get("mogheiyat_varagh") or "").strip()[:200] or None
        if can_edit("tedad_baste"):
            item.tedad_baste = parse_int("tedad_baste", "تعداد بسته")
        if can_edit("vazn_esnadi"):
            item.vazn_esnadi = parse_float("vazn_esnadi", "وزن اسنادی")
        if can_edit("tarikh_vorud"):
            item.tarikh_vorud = parse_date("tarikh_vorud")
        if can_edit("saat_vorud"):
            item.saat_vorud = parse_time("saat_vorud")
        if can_edit("tarikh_khoruj"):
            item.tarikh_khoruj = parse_date("tarikh_khoruj")
        if can_edit("saat_khoruj"):
            item.saat_khoruj = parse_time("saat_khoruj")
        if can_edit("masoul_sabt"):
            item.masoul_sabt = (request.form.get("masoul_sabt") or "").strip()[:100] or None
        if can_edit("vazn_por"):
            item.vazn_por = parse_float("vazn_por", "وزن پُر باسکول")
        if can_edit("vazn_khali"):
            item.vazn_khali = parse_float("vazn_khali", "وزن خالی باسکول")

        if item.vazn_por is not None and item.vazn_khali is not None and item.vazn_khali > item.vazn_por:
            raise ValueError("وزن خالی نمی‌تواند از وزن پُر بیشتر باشد.")

        if can_edit("ghabz_por_file"):
            item.ghabz_por_file = save_receipt("ghabz_por_file", item.ghabz_por_file)
        if can_edit("ghabz_khali_file"):
            item.ghabz_khali_file = save_receipt("ghabz_khali_file", item.ghabz_khali_file)
        item.ghabz_por_piost = "بله" if item.ghabz_por_file else "خیر"
        item.ghabz_khali_piost = "بله" if item.ghabz_khali_file else "خیر"

        if can_edit("vaziyat_havaleh"):
            status = (request.form.get("vaziyat_havaleh") or "ثبت‌شده").strip()
            if status not in STATUSES:
                raise ValueError("وضعیت حواله معتبر نیست.")
            if status == "تکمیل‌شده" and current_user.role not in (ROLE_ADMIN, ROLE_PROJECT):
                raise ValueError("تکمیل نهایی حواله فقط برای مدیر یا نقش پروژه مجاز است.")
            item.vaziyat_havaleh = status

        if can_edit("vaziyat_moghayerat"):
            discrepancy = (request.form.get("vaziyat_moghayerat") or "بدون مغایرت").strip()
            if discrepancy not in MOGHAYERAT:
                raise ValueError("وضعیت مغایرت معتبر نیست.")
            item.vaziyat_moghayerat = discrepancy

        if can_edit("sharh_moghayerat"):
            item.sharh_moghayerat = (request.form.get("sharh_moghayerat") or "").strip() or None
        if can_edit("masoul_peygiri"):
            item.masoul_peygiri = (request.form.get("masoul_peygiri") or "").strip()[:100] or None
        if can_edit("tarikh_peygiri"):
            item.tarikh_peygiri = parse_date("tarikh_peygiri")
        if can_edit("natije_peygiri"):
            item.natije_peygiri = (request.form.get("natije_peygiri") or "").strip() or None
        if can_edit("taeed_nahayi"):
            item.taeed_nahayi = (request.form.get("taeed_nahayi") or "").strip()[:100] or None
        if can_edit("tarikh_taeed"):
            item.tarikh_taeed = parse_date("tarikh_taeed")
        if can_edit("tozihat"):
            item.tozihat = (request.form.get("tozihat") or "").strip() or None


        # ----- چندردیفی کردن کالا با سازگاری با حواله‌های قدیمی -----
        row_fields = ["item_code_kala", "item_sharh_kala", "item_moshakhasat_varagh", "item_mogheiyat_varagh", "item_tedad_baste", "item_vazn_esnadi"]
        row_values = [request.form.getlist(field) for field in row_fields]
        submitted_rows = []
        for idx in range(max([len(v) for v in row_values] or [0])):
            vals = [(row_values[k][idx] if idx < len(row_values[k]) else "").strip() for k in range(6)]
            if not any(vals): continue
            code, name, spec, location, qty_raw, weight_raw = vals
            if not code: raise ValueError(f"کد کالا در ردیف {idx + 1} انتخاب نشده است.")
            try: qty = int(qty_raw) if qty_raw else None
            except ValueError: raise ValueError(f"تعداد بسته در ردیف {idx + 1} باید عدد صحیح باشد.")
            if qty is not None and (qty < 0 or qty > 10_000_000): raise ValueError(f"تعداد بسته در ردیف {idx + 1} معتبر نیست.")
            try: weight = float(weight_raw.replace(",", "")) if weight_raw else None
            except ValueError: raise ValueError(f"وزن اسنادی ردیف {idx + 1} باید عددی باشد.")
            if weight is not None and (not math.isfinite(weight) or weight < 0): raise ValueError(f"وزن اسنادی ردیف {idx + 1} معتبر نیست.")
            submitted_rows.append({"radif":len(submitted_rows)+1,"code_kala":code[:30],"sharh_kala":name[:100] or None,
                "moshakhasat_varagh":spec[:200] or None,"mogheiyat_varagh":location[:200] or None,"tedad_baste":qty,"vazn_esnadi":weight})
        has_multi_rows = bool(submitted_rows)
        # نقش‌هایی که اجازه تغییر کد/وزن/تعداد کالا را ندارند، از طریق فیلدهای تکرارشونده هم نمی‌توانند این محدودیت را دور بزنند.
        if has_multi_rows and any(not can_edit(field) for field in ("code_kala", "vazn_esnadi", "tedad_baste")):
            old_rows = [(r.code_kala or "", r.sharh_kala or "", r.moshakhasat_varagh or "",
                         r.mogheiyat_varagh or "", r.tedad_baste, r.vazn_esnadi) for r in item.kala_items]
            if not old_rows:
                old_rows = [(item.code_kala or "", item.sharh_kala or "", item.moshakhasat_varagh or "",
                             item.mogheiyat_varagh or "", item.tedad_baste, item.vazn_esnadi)]
            new_rows = [(r["code_kala"], r["sharh_kala"] or "", r["moshakhasat_varagh"] or "",
                         r["mogheiyat_varagh"] or "", r["tedad_baste"], r["vazn_esnadi"]) for r in submitted_rows]
            if not item.id or new_rows != old_rows:
                raise ValueError("نقش کاربری شما اجازه افزودن یا تغییر ردیف‌های کالا، تعداد و وزن اسنادی را ندارد.")
        if has_multi_rows:
            if item.inventory_deducted:
                existing = [(r.code_kala or "",r.sharh_kala or "",r.moshakhasat_varagh or "",r.mogheiyat_varagh or "",r.tedad_baste,r.vazn_esnadi) for r in item.kala_items]
                if not existing: existing=[(item.code_kala or "",item.sharh_kala or "",item.moshakhasat_varagh or "",item.mogheiyat_varagh or "",item.tedad_baste,item.vazn_esnadi)]
                requested=[(r["code_kala"],r["sharh_kala"] or "",r["moshakhasat_varagh"] or "",r["mogheiyat_varagh"] or "",r["tedad_baste"],r["vazn_esnadi"]) for r in submitted_rows]
                if requested != existing: raise ValueError("ردیف‌های کالای حواله‌ای که از موجودی آن کسر شده قابل تغییر نیستند؛ برای اصلاح با مدیر سامانه هماهنگ کنید.")
            else:
                item.kala_items.clear()
                for row in submitted_rows: item.kala_items.append(HavalehItem(**row))
                first=submitted_rows[0]
                item.code_kala,item.sharh_kala=first["code_kala"],first["sharh_kala"]
                item.moshakhasat_varagh,item.mogheiyat_varagh=first["moshakhasat_varagh"],first["mogheiyat_varagh"]
                item.tedad_baste=first["tedad_baste"]
                item.vazn_esnadi=round(sum(r["vazn_esnadi"] or 0 for r in submitted_rows),3)
        else:
            submitted_rows=[{"radif":1,"code_kala":item.code_kala,"sharh_kala":item.sharh_kala,"moshakhasat_varagh":item.moshakhasat_varagh,
                "mogheiyat_varagh":item.mogheiyat_varagh,"tedad_baste":item.tedad_baste,"vazn_esnadi":item.vazn_esnadi}]
        # ابتدا همه ردیف‌ها را کنترل کن، سپس کسر موجودی انجام شود.
        if item.vaziyat_havaleh == "تکمیل‌شده" and not item.inventory_deducted:
            if item.vazn_khales is None: raise ValueError("برای تکمیل حواله، وزن پُر و خالی باسکول را ثبت کنید.")
            if has_multi_rows:
                if any(not r["code_kala"] or r["vazn_esnadi"] is None or r["vazn_esnadi"] <= 0 for r in submitted_rows):
                    raise ValueError("برای تکمیل حواله، کد و وزن اسنادی مثبت هر ردیف کالا الزامی است.")
                totals, stocks, deductions = {}, {}, []
                for row in submitted_rows:
                    candidates=Inventory.query.filter_by(code_kala=row["code_kala"]).order_by(Inventory.id).all()
                    if row["mogheiyat_varagh"]: candidates=[x for x in candidates if (x.mogheiyat or "").strip()==row["mogheiyat_varagh"]]
                    if len(candidates)>1:
                        spec=(row["moshakhasat_varagh"] or "").strip()
                        matches=[x for x in candidates if spec and spec in {(x.moshakhasat or "").strip(),(x.name_kala or "").strip()}]
                        if len(matches)!=1: raise ValueError(f"ردیف {row['radif']}: کد کالا چند موجودی دارد؛ مشخصات و محل را مطابق ایندکس انتخاب کنید.")
                        stock=matches[0]
                    elif len(candidates)==1: stock=candidates[0]
                    else: raise ValueError(f"ردیف {row['radif']}: کد کالا در ایندکس موجودی پیدا نشد.")
                    w=row["vazn_esnadi"]; deductions.append((stock,w,row)); totals[stock.id]=totals.get(stock.id,0)+w; stocks[stock.id]=stock
                for sid,total in totals.items():
                    stock=stocks[sid]
                    if stock.mojudi_vaghei is None or stock.mojudi_vaghei<total: raise ValueError("موجودی یکی از کالاها برای جمع ردیف‌های این حواله کافی نیست؛ هیچ کسر انجام نشد.")
                for sid,total in totals.items(): stocks[sid].mojudi_vaghei=round(stocks[sid].mojudi_vaghei-total,3)
                for child in item.kala_items: child.inventory_deducted=True
                item.inventory_deducted=True
                log_action("inventory_deducted","Havaleh",item.id,"کسر چندردیفی: "+"; ".join(f"{r['code_kala']} وزن اسنادی={w}" for _,w,r in deductions))
            else:
                if not item.code_kala: raise ValueError("برای تکمیل حواله، کد کالای موجودی باید مشخص شود.")
                candidates=Inventory.query.filter_by(code_kala=item.code_kala).order_by(Inventory.id).all()
                if not candidates: raise ValueError("کد کالا در ایندکس موجودی پیدا نشد؛ ابتدا موجودی را بررسی کنید.")
                if item.mogheiyat_varagh: candidates=[row for row in candidates if (row.mogheiyat or "").strip()==item.mogheiyat_varagh]
                if len(candidates)>1:
                    spec=(item.moshakhasat_varagh or "").strip()
                    matches=[row for row in candidates if spec and spec in {(row.moshakhasat or "").strip(),(row.name_kala or "").strip()}]
                    if len(matches)!=1: raise ValueError("این کد کالا چند ردیف موجودی دارد؛ مشخصات ورق را دقیقاً مطابق ایندکس انتخاب کنید.")
                    inventory=matches[0]
                elif len(candidates)==1: inventory=candidates[0]
                else: raise ValueError("محل انتخاب‌شده برای این کد کالا در ایندکس موجودی پیدا نشد.")
                if inventory.mojudi_vaghei is None or inventory.mojudi_vaghei<item.vazn_khales: raise ValueError("موجودی کافی نیست یا موجودی واقعی تعریف نشده است؛ کسر انجام نشد.")
                inventory.mojudi_vaghei=round(inventory.mojudi_vaghei-item.vazn_khales,3); item.inventory_deducted=True
                log_action("inventory_deducted","Havaleh",item.id,f"کد کالا={item.code_kala}; وزن خالص={item.vazn_khales}; موجودی پس از کسر={inventory.mojudi_vaghei}")

        item.updated_by_id = current_user.id
        return old_status

    # ---------- auth ----------
    @auth_bp.route("/login", methods=["GET", "POST"])
    def login():
        if current_user.is_authenticated:
            return redirect(url_for("dashboard.index"))

        if request.method == "POST":
            username = (request.form.get("username") or "").strip()
            password = request.form.get("password") or ""

            try:
                user = User.query.filter_by(username=username).first()
            except OperationalError:
                db.session.rollback()
                db.session.remove()
                try:
                    user = User.query.filter_by(username=username).first()
                except OperationalError:
                    db.session.rollback()
                    db.session.remove()
                    app.logger.exception("Database unavailable while checking login")
                    abort(
                        503,
                        description="ارتباط با پایگاه داده موقتاً برقرار نیست؛ "
                                    "چند لحظه بعد دوباره تلاش کنید.",
                    )

            if user and user.is_active_user and user.check_password(password):
                session.clear()
                login_user(user, remember=False, fresh=True)
                session["_csrf_token"] = secrets.token_urlsafe(32)
                flash("با موفقیت وارد شدید.", "success")
                target = _safe_next_url(request.args.get("next")) or url_for("dashboard.index")
                return redirect(target)

            # Timing-safe fallback for missing users
            if not user:
                try:
                    from werkzeug.security import generate_password_hash
                    generate_password_hash("timing-safe-dummy")
                except Exception:
                    pass

            flash("نام کاربری یا رمز عبور صحیح نیست.", "danger")

        return render_template("login.html")

    @auth_bp.route("/logout", methods=["POST"])
    @login_required
    def logout():
        logout_user()
        session.clear()
        flash("از سامانه خارج شدید.", "info")
        return redirect(url_for("auth.login"))

    # ---------- dashboard ----------
    @dashboard_bp.route("/")
    @login_required
    def index():
        all_items = Havaleh.query.order_by(Havaleh.created_at.desc()).all()
        total = len(all_items)
        completed = sum(h.vaziyat_havaleh == "تکمیل‌شده" for h in all_items)
        incomplete = sum(h.vaziyat_havaleh == "ناقص" for h in all_items)
        need_review = sum(
            h.vaziyat_havaleh == "نیازمند بررسی"
            or h.vaziyat_moghayerat != "بدون مغایرت"
            for h in all_items
        )
        sum_esnadi = sum(h.vazn_esnadi or 0 for h in all_items)
        sum_khales = sum(h.vazn_khales or 0 for h in all_items)
        diffs = [h.ekhtelaf_vazn for h in all_items if h.ekhtelaf_vazn is not None]
        durations = [
            h.moddat_tavaghof_saat
            for h in all_items
            if h.moddat_tavaghof_saat is not None and h.moddat_tavaghof_saat >= 0
        ]
        stats = {
            "total": total, "completed": completed,
            "incomplete": incomplete, "need_review": need_review,
            "sum_esnadi": sum_esnadi, "sum_khales": sum_khales,
            "balance": sum_khales - sum_esnadi,
            "pos_diff": sum(d > 0 for d in diffs),
            "neg_diff": sum(d < 0 for d in diffs),
            "sum_pos": sum(d for d in diffs if d > 0),
            "sum_neg": sum(d for d in diffs if d < 0),
            "no_ghabz_por": sum(not h.ghabz_por_file for h in all_items),
            "no_ghabz_khali": sum(not h.ghabz_khali_file for h in all_items),
            "total_duration": round(sum(durations), 2),
            "avg_duration": round(sum(durations) / len(durations), 2) if durations else 0,
        }
        followups = [
            h for h in all_items
            if h.vaziyat_havaleh in ("ناقص", "نیازمند بررسی")
            or h.vaziyat_moghayerat != "بدون مغایرت"
            or not h.ghabz_por_file
            or not h.ghabz_khali_file
        ]
        return render_template("index.html", stats=stats, followups=followups[:100])

    # ---------- havaleh ----------
    @havaleh_bp.route("/havaleh/blank-print")
    @login_required
    def blank_print():
        return render_template("havaleh_blank_print.html")

    @havaleh_bp.route("/havaleh")
    @login_required
    def havaleh_list():
        q = (request.args.get("q") or "").strip()[:100]
        status = request.args.get("status", "")
        query = Havaleh.query
        if q:
            like = f"%{q}%"
            query = query.filter(or_(
                Havaleh.shomare_havaleh.ilike(like),
                Havaleh.pelak.ilike(like),
                Havaleh.name_ranande.ilike(like),
                Havaleh.noe_vasileh.ilike(like),
                Havaleh.sharh_kala.ilike(like),
                Havaleh.code_kala.ilike(like),
                Havaleh.moshakhasat_varagh.ilike(like),
            ))
        if status in STATUSES:
            query = query.filter_by(vaziyat_havaleh=status)
        items = query.order_by(Havaleh.created_at.desc()).limit(1000).all()
        return render_template(
            "havaleh_list.html", items=items, q=q, status=status, statuses=STATUSES,
        )

    @havaleh_bp.route("/havaleh/new", methods=["GET", "POST"])
    @require_roles(ROLE_ADMIN, ROLE_HERASAT, ROLE_ANBAR, ROLE_PROJECT)
    def havaleh_new():
        item = Havaleh()
        if request.method == "POST":
            try:
                supplied_number = (request.form.get("shomare_havaleh") or "").strip()
                item.shomare_havaleh = supplied_number or next_voucher_number()
                if Havaleh.query.filter_by(shomare_havaleh=item.shomare_havaleh).first():
                    raise ValueError("این شماره حواله قبلاً ثبت شده است.")
                item.created_by_id = current_user.id
                db.session.add(item)
                db.session.flush()
                apply_form(
                    item, creating=True,
                    editable_fields=role_editable_fields(current_user.role),
                )
                log_action("create", "Havaleh", item.id, f"شماره حواله={item.shomare_havaleh}")
                db.session.commit()
                flash(f"حواله {item.shomare_havaleh} ثبت شد.", "success")
                return redirect(url_for("havaleh.havaleh_list"))
            except ValueError as exc:
                db.session.rollback()
                flash(str(exc), "danger")
            except IntegrityError:
                db.session.rollback()
                flash("شماره حواله تکراری است؛ دوباره تلاش کنید.", "danger")
            except OperationalError:
                db.session.rollback()
                app.logger.exception("Database error on havaleh_new")
                flash("ارتباط با پایگاه داده موقتاً برقرار نیست؛ دوباره تلاش کنید.", "danger")

        inventory = Inventory.query.filter(Inventory.mogheiyat.ilike("%لیدان%"), ~Inventory.mogheiyat.ilike("%ساب%")).order_by(Inventory.code_kala).all()
        return render_template(
            "havaleh_form.html", item=None, inventory=inventory, statuses=STATUSES,
            moghayerat_list=MOGHAYERAT, bale_kheyr=BALE_KHEYR,
        )

    @havaleh_bp.route("/havaleh/<int:id>/edit", methods=["GET", "POST"])
    @require_roles(ROLE_ADMIN, ROLE_HERASAT, ROLE_ANBAR, ROLE_PROJECT)
    def havaleh_edit(id):
        item = db.get_or_404(Havaleh, id)
        if request.method == "POST":
            try:
                old_status = item.vaziyat_havaleh
                editable = role_editable_fields(current_user.role)
                enforce_role_field_permissions(item, editable)
                apply_form(item, editable_fields=editable)
                log_action(
                    "update", "Havaleh", item.id,
                    f"وضعیت: {old_status} ← {item.vaziyat_havaleh}",
                )
                db.session.commit()
                flash("تغییرات حواله ذخیره شد.", "success")
                return redirect(url_for("havaleh.havaleh_list"))
            except ValueError as exc:
                db.session.rollback()
                flash(str(exc), "danger")
            except IntegrityError:
                db.session.rollback()
                flash("ذخیره تغییرات ناموفق بود؛ دوباره تلاش کنید.", "danger")
            except OperationalError:
                db.session.rollback()
                app.logger.exception("Database error on havaleh_edit")
                flash("ارتباط با پایگاه داده موقتاً برقرار نیست؛ دوباره تلاش کنید.", "danger")

        inventory = Inventory.query.filter(Inventory.mogheiyat.ilike("%لیدان%"), ~Inventory.mogheiyat.ilike("%ساب%")).order_by(Inventory.code_kala).all()
        return render_template(
            "havaleh_form.html", item=item, inventory=inventory, statuses=STATUSES,
            moghayerat_list=MOGHAYERAT, bale_kheyr=BALE_KHEYR,
        )

    @havaleh_bp.route("/havaleh/<int:id>/delete", methods=["POST"])
    @require_roles(ROLE_ADMIN)
    def havaleh_delete(id):
        item = db.get_or_404(Havaleh, id)
        if item.inventory_deducted or item.vaziyat_havaleh == "تکمیل‌شده":
            flash(
                "حواله‌ای که روی موجودی اثر گذاشته یا تکمیل شده، حذف نمی‌شود؛ "
                "برای اصلاح، با مدیر سامانه هماهنگ کنید.",
                "danger",
            )
            return redirect(url_for("havaleh.havaleh_list"))
        try:
            log_action("delete", "Havaleh", item.id, f"شماره حواله={item.shomare_havaleh}")
            db.session.delete(item)
            db.session.commit()
            flash("حواله حذف شد.", "success")
        except OperationalError:
            db.session.rollback()
            flash("ارتباط با پایگاه داده موقتاً برقرار نیست؛ دوباره تلاش کنید.", "danger")
        return redirect(url_for("havaleh.havaleh_list"))

    # ---------- users ----------
    @users_bp.route("/users", methods=["GET", "POST"])
    @require_roles(ROLE_ADMIN)
    def users():
        if request.method == "POST":
            username = (request.form.get("username") or "").strip()
            full_name = (request.form.get("full_name") or "").strip()[:120] or None
            password = request.form.get("password") or ""
            password_confirm = request.form.get("password_confirm") or password
            role = (request.form.get("role") or ROLE_VIEWER).strip()

            if (
                not username
                or len(username) > 80
                or not username.replace("_", "").replace("-", "").isalnum()
            ):
                flash("نام کاربری باید فقط شامل حروف انگلیسی، عدد، خط تیره یا زیرخط باشد.", "danger")
            elif len(password) < 12:
                flash("رمز عبور باید حداقل ۱۲ نویسه باشد.", "danger")
            elif len(password) > 128:
                flash("رمز عبور نباید بیش از ۱۲۸ نویسه باشد.", "danger")
            elif password != password_confirm:
                flash("رمز عبور و تکرار آن یکسان نیستند.", "danger")
            elif role not in (ROLE_ADMIN, ROLE_HERASAT, ROLE_ANBAR, ROLE_PROJECT, ROLE_VIEWER):
                flash("نقش انتخاب‌شده معتبر نیست.", "danger")
            elif User.query.filter_by(username=username).first():
                flash("این نام کاربری قبلاً ثبت شده است.", "danger")
            else:
                try:
                    user = User(
                        username=username, full_name=full_name,
                        role=role, is_active_user=True,
                    )
                    user.set_password(password)
                    db.session.add(user)
                    db.session.flush()
                    log_action(
                        "user_created", "User", user.id,
                        f"username={username}; role={role}",
                    )
                    db.session.commit()
                    flash("کاربر جدید ساخته شد.", "success")
                    return redirect(url_for("users.users"))
                except OperationalError:
                    db.session.rollback()
                    flash("ارتباط با پایگاه داده موقتاً برقرار نیست؛ دوباره تلاش کنید.", "danger")

        users_list = User.query.order_by(User.username).all()
        return render_template(
            "users.html", users=users_list,
            roles={
                ROLE_ADMIN: "مدیر سیستم", ROLE_HERASAT: "حراست",
                ROLE_ANBAR: "انبار", ROLE_PROJECT: "پروژه", ROLE_VIEWER: "مشاهده‌گر",
            },
        )

    @users_bp.route("/users/<int:id>/toggle", methods=["POST"])
    @require_roles(ROLE_ADMIN)
    def user_toggle(id):
        user = db.get_or_404(User, id)
        if user.id == current_user.id:
            flash("امکان غیرفعال‌کردن حساب خودتان وجود ندارد.", "danger")
        else:
            try:
                user.is_active_user = not user.is_active_user
                log_action(
                    "user_enabled" if user.is_active_user else "user_disabled",
                    "User", user.id, f"username={user.username}",
                )
                db.session.commit()
                flash("وضعیت کاربر تغییر کرد.", "success")
            except OperationalError:
                db.session.rollback()
                flash("ارتباط با پایگاه داده موقتاً برقرار نیست؛ دوباره تلاش کنید.", "danger")
        return redirect(url_for("users.users"))

    # ---------- inventory ----------
    @inventory_bp.route("/inventory/template.csv")
    @require_roles(ROLE_ADMIN)
    def inventory_template():
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "code_kala", "name_kala", "zekhamat", "moshakhasat", "vahed",
            "mojudi_vaghei", "tedad_sheet", "vazn_har_sheet", "vazn_goruh",
            "mogheiyat", "radif",
        ])
        response = make_response("\ufeff" + output.getvalue())
        response.headers["Content-Type"] = "text/csv; charset=utf-8"
        response.headers["Content-Disposition"] = (
            "attachment; filename=varagh-inventory-template.csv"
        )
        return response

    @inventory_bp.route("/inventory/import", methods=["POST"])
    @require_roles(ROLE_ADMIN)
    def inventory_import():
        upload = request.files.get("inventory_file")
        if not upload or not upload.filename:
            flash("فایل CSV یا XLSX موجودی را انتخاب کنید.", "danger")
            return redirect(url_for("inventory.inventory_list"))

        filename = secure_filename(upload.filename)
        extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if extension not in ("csv", "xlsx"):
            flash("فقط فایل CSV یا XLSX پذیرفته می‌شود.", "danger")
            return redirect(url_for("inventory.inventory_list"))

        try:
            if extension == "csv":
                rows = list(csv.DictReader(
                    io.StringIO(upload.read().decode("utf-8-sig"))
                ))
            else:
                from openpyxl import load_workbook
                workbook = load_workbook(
                    io.BytesIO(upload.read()), read_only=True, data_only=True,
                )
                sheet = workbook.active
                values = sheet.iter_rows(values_only=True)
                header = next(values, None)
                if not header:
                    raise ValueError("فایل اکسل خالی است.")
                keys = [str(value or "").strip() for value in header]
                rows = [
                    dict(zip(keys, row))
                    for row in values
                    if any(value is not None for value in row)
                ]
                workbook.close()

            if not rows:
                raise ValueError("فایل هیچ ردیف داده‌ای ندارد.")
            if len(rows) > 5000:
                raise ValueError("در هر بار حداکثر ۵۰۰۰ ردیف قابل ورود است.")

            required = {"code_kala", "mojudi_vaghei"}
            headers = {
                str(key).strip()
                for row in rows
                for key in row.keys()
                if key is not None
            }
            if not required.issubset(headers):
                raise ValueError(
                    "ستون‌های code_kala و mojudi_vaghei الزامی هستند؛ "
                    "از فایل الگو استفاده کنید."
                )

            def cell_text(value):
                return str(value).strip() if value is not None else ""

            def cell_number(value, field, row_number, required=False):
                raw = cell_text(value).replace(",", "")
                if not raw:
                    if required:
                        raise ValueError(
                            f"ردیف {row_number}: مقدار {field} الزامی است."
                        )
                    return None
                try:
                    number = float(raw)
                except (TypeError, ValueError):
                    raise ValueError(
                        f"ردیف {row_number}: مقدار {field} عدد معتبر نیست."
                    )
                if not math.isfinite(number):
                    raise ValueError(
                        f"ردیف {row_number}: مقدار {field} نامعتبر است."
                    )
                if number < 0:
                    raise ValueError(
                        f"ردیف {row_number}: مقدار {field} نمی‌تواند منفی باشد."
                    )
                return number

            prepared = []
            seen = set()
            for row_number, row in enumerate(rows, start=2):
                record = {str(k).strip(): v for k, v in row.items() if k is not None}
                code = cell_text(record.get("code_kala"))[:30]
                if not code:
                    raise ValueError(f"ردیف {row_number}: کد کالا الزامی است.")
                spec = cell_text(record.get("moshakhasat"))[:200]
                location = cell_text(record.get("mogheiyat"))[:200]
                identity = (code, spec, location)
                if identity in seen:
                    raise ValueError(
                        f"ردیف {row_number}: ترکیب کد کالا، مشخصات و موقعیت تکراری است."
                    )
                seen.add(identity)
                prepared.append({
                    "code_kala": code,
                    "name_kala": cell_text(record.get("name_kala"))[:100] or None,
                    "zekhamat": cell_number(record.get("zekhamat"), "zekhamat", row_number),
                    "moshakhasat": spec or None,
                    "vahed": cell_text(record.get("vahed"))[:20] or "کیلو گرم",
                    "mojudi_vaghei": cell_number(
                        record.get("mojudi_vaghei"), "mojudi_vaghei",
                        row_number, True,
                    ),
                    "tedad_sheet": cell_number(
                        record.get("tedad_sheet"), "tedad_sheet", row_number,
                    ),
                    "vazn_har_sheet": cell_number(
                        record.get("vazn_har_sheet"), "vazn_har_sheet", row_number,
                    ),
                    "vazn_goruh": cell_number(
                        record.get("vazn_goruh"), "vazn_goruh", row_number,
                    ),
                    "mogheiyat": location or None,
                    "radif": (
                        int(cell_number(record.get("radif"), "radif", row_number))
                        if cell_text(record.get("radif")) else None
                    ),
                })

            changed = 0
            for record in prepared:
                query = Inventory.query.filter_by(
                    code_kala=record["code_kala"],
                    moshakhasat=record["moshakhasat"],
                    mogheiyat=record["mogheiyat"],
                )
                item = query.first()
                if item is None:
                    item = Inventory(**record)
                    db.session.add(item)
                else:
                    for key, value in record.items():
                        setattr(item, key, value)
                changed += 1
            log_action(
                "inventory_import", "Inventory", None,
                f"ردیف‌های ثبت/به‌روزرسانی‌شده={changed}",
            )
            db.session.commit()
            flash(
                f"ورود موجودی انجام شد؛ {changed} ردیف ثبت یا به‌روزرسانی شد.",
                "success",
            )
        except (ValueError, UnicodeDecodeError, csv.Error) as exc:
            db.session.rollback()
            flash(str(exc), "danger")
        except OperationalError:
            db.session.rollback()
            app.logger.exception("Database unavailable during inventory import")
            flash("ارتباط با پایگاه داده موقتاً برقرار نیست؛ دوباره تلاش کنید.", "danger")
        except Exception:
            db.session.rollback()
            app.logger.exception("Inventory import failed")
            flash("ورود فایل ناموفق بود؛ ساختار فایل را بررسی کنید.", "danger")

        return redirect(url_for("inventory.inventory_list"))

    @inventory_bp.route("/inventory")
    @login_required
    def inventory_list():
        items = Inventory.query.filter(Inventory.mogheiyat.ilike("%لیدان%"), ~Inventory.mogheiyat.ilike("%ساب%")).order_by(Inventory.code_kala, Inventory.id).all()
        total_kg = sum((i.mojudi_vaghei or 0) for i in items)
        unique_count = len({i.code_kala for i in items})
        return render_template(
            "inventory.html", items=items,
            total_kg=total_kg, unique_count=unique_count,
        )

    # ---------- reports ----------
    @reports_bp.route("/export/excel")
    @login_required
    def export_excel():
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            abort(500, description="کتابخانه خروجی اکسل نصب نشده است.")

        wb = Workbook()
        ws = wb.active
        ws.title = "حواله‌ها"
        headers = [
            "شماره حواله", "تاریخ حواله", "نوع وسیله", "پلاک کشنده", "نام راننده",
            "شماره تماس راننده", "شرح کالا", "مشخصات ورق", "کد کالا", "محل موجودی",
            "تعداد بسته", "وزن اسنادی", "وزن پر", "وزن خالی", "وزن خالص",
            "اختلاف وزن", "تاریخ ورود", "ساعت ورود", "تاریخ خروج", "ساعت خروج",
            "وضعیت", "مغایرت", "قبض پر", "قبض خالی", "توضیحات",
        ]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")

        for h in Havaleh.query.order_by(Havaleh.created_at.desc()).all():
            ws.append([
                h.shomare_havaleh,
                h.tarikh_havaleh.isoformat() if h.tarikh_havaleh else "",
                h.noe_vasileh, h.pelak, h.name_ranande, h.mobile_ranande,
                h.sharh_kala, h.moshakhasat_varagh, h.code_kala,
                h.mogheiyat_varagh, h.tedad_baste, h.vazn_esnadi,
                h.vazn_por, h.vazn_khali, h.vazn_khales, h.ekhtelaf_vazn,
                h.tarikh_vorud.isoformat() if h.tarikh_vorud else "",
                h.saat_vorud.strftime("%H:%M") if h.saat_vorud else "",
                h.tarikh_khoruj.isoformat() if h.tarikh_khoruj else "",
                h.saat_khoruj.strftime("%H:%M") if h.saat_khoruj else "",
                h.vaziyat_havaleh, h.vaziyat_moghayerat,
                h.ghabz_por_piost, h.ghabz_khali_piost, h.tozihat,
            ])

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            width = min(max(max(len(str(c.value or "")) for c in col) + 2, 12), 32)
            ws.column_dimensions[col[0].column_letter].width = width

        output = io.BytesIO()
        wb.save(output)
        response = make_response(output.getvalue())
        response.headers["Content-Type"] = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response.headers["Content-Disposition"] = (
            "attachment; filename=varagh-havaleh.xlsx"
        )
        return response

    @reports_bp.route("/reconciliation", methods=["GET", "POST"])
    @require_roles(ROLE_ADMIN, ROLE_PROJECT)
    def reconciliation():
        control = db.session.get(TransferControl, 1)
        if control is None:
            control = TransferControl(id=1)
            db.session.add(control)
            db.session.commit()

        if request.method == "POST":
            action = (request.form.get("action") or "").strip()
            try:
                if action == "save_valuation":
                    control.agreement_signed_on = parse_date("agreement_signed_on")
                    control.valuation_announced_on = parse_date("valuation_announced_on")
                    control.valuation_price_per_kg = parse_float(
                        "valuation_price_per_kg", "قیمت کارشناسی هر کیلوگرم",
                    )
                    control.valuation_reference = (
                        (request.form.get("valuation_reference") or "").strip()[:300]
                        or None
                    )
                    if (
                        control.valuation_announced_on
                        and control.valuation_price_per_kg is None
                    ):
                        raise ValueError(
                            "برای ثبت تاریخ اعلام قیمت، مبلغ کارشناسی هر کیلوگرم را "
                            "نیز وارد کنید."
                        )
                    log_action(
                        "valuation_updated", "TransferControl", 1,
                        f"تاریخ اعلام={control.valuation_announced_on}; "
                        f"قیمت هر کیلوگرم={control.valuation_price_per_kg}",
                    )
                    db.session.commit()
                    flash("اطلاعات کارشناسی و تاریخ‌های مربوط ذخیره شد.", "success")

                elif action == "finalize_summary":
                    items = Havaleh.query.order_by(Havaleh.shomare_havaleh).all()
                    durations = [
                        h.moddat_tavaghof_saat for h in items
                        if h.moddat_tavaghof_saat is not None
                        and h.moddat_tavaghof_saat >= 0
                    ]
                    snapshot = {
                        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                        "voucher_count": len(items),
                        "completed_count": sum(
                            h.vaziyat_havaleh == "تکمیل‌شده" for h in items
                        ),
                        "document_weight_kg": round(
                            sum(h.vazn_esnadi or 0 for h in items), 2,
                        ),
                        "net_scale_weight_kg": round(
                            sum(
                                h.vazn_khales or 0 for h in items
                                if h.vazn_khales is not None
                            ), 2,
                        ),
                        "weight_difference_kg": round(
                            sum(h.ekhtelaf_vazn or 0 for h in items), 2,
                        ),
                        "duration_hours": round(sum(durations), 2),
                        "missing_full_receipt": sum(
                            not h.ghabz_por_file for h in items
                        ),
                        "missing_empty_receipt": sum(
                            not h.ghabz_khali_file for h in items
                        ),
                        "incomplete_or_review": sum(
                            h.vaziyat_havaleh in ("ناقص", "نیازمند بررسی")
                            or h.vaziyat_moghayerat != "بدون مغایرت"
                            or h.vazn_khales is None
                            for h in items
                        ),
                        "voucher_numbers": [h.shomare_havaleh for h in items],
                    }
                    control.final_summary_signed_on = parse_date("final_summary_signed_on")
                    control.final_summary_signed_by = (
                        (request.form.get("final_summary_signed_by") or "").strip()[:200]
                        or None
                    )
                    control.final_summary_notes = (
                        (request.form.get("final_summary_notes") or "").strip()[:2000]
                        or None
                    )
                    if (
                        not control.final_summary_signed_on
                        or not control.final_summary_signed_by
                    ):
                        raise ValueError(
                            "برای ثبت صورتجلسه نهایی، تاریخ و نام/سمت تأییدکننده را "
                            "وارد کنید."
                        )
                    control.final_summary_snapshot = json.dumps(
                        snapshot, ensure_ascii=False,
                    )
                    log_action(
                        "final_reconciliation_saved", "TransferControl", 1,
                        f"تاریخ={control.final_summary_signed_on}; "
                        f"تأییدکننده={control.final_summary_signed_by}; "
                        f"تعداد حواله={len(items)}",
                    )
                    db.session.commit()
                    flash(
                        "صورتجلسه نهایی با snapshot مقادیر فعلی ثبت شد. "
                        "برای تغییرات بعدی، گزارش جدید را جداگانه ثبت کنید.",
                        "success",
                    )
                else:
                    raise ValueError("عملیات درخواستی معتبر نیست.")

            except ValueError as exc:
                db.session.rollback()
                flash(str(exc), "danger")
            except IntegrityError:
                db.session.rollback()
                flash("ذخیره اطلاعات ناموفق بود؛ دوباره تلاش کنید.", "danger")
            except OperationalError:
                db.session.rollback()
                app.logger.exception("Database error on reconciliation")
                flash("ارتباط با پایگاه داده موقتاً برقرار نیست؛ دوباره تلاش کنید.", "danger")

            return redirect(url_for("reports.reconciliation"))

        items = Havaleh.query.order_by(Havaleh.created_at.desc()).all()
        durations = [
            h.moddat_tavaghof_saat for h in items
            if h.moddat_tavaghof_saat is not None and h.moddat_tavaghof_saat >= 0
        ]
        stats = {
            "total": len(items),
            "completed": sum(h.vaziyat_havaleh == "تکمیل‌شده" for h in items),
            "incomplete": sum(h.vaziyat_havaleh == "ناقص" for h in items),
            "review": sum(
                h.vaziyat_havaleh == "نیازمند بررسی"
                or h.vaziyat_moghayerat != "بدون مغایرت"
                for h in items
            ),
            "document_weight": round(sum(h.vazn_esnadi or 0 for h in items), 2),
            "net_weight": round(
                sum(h.vazn_khales or 0 for h in items if h.vazn_khales is not None), 2,
            ),
            "difference": round(sum(h.ekhtelaf_vazn or 0 for h in items), 2),
            "duration": round(sum(durations), 2),
            "missing_full": sum(not h.ghabz_por_file for h in items),
            "missing_empty": sum(not h.ghabz_khali_file for h in items),
            "missing_times": sum(
                not all([h.tarikh_vorud, h.saat_vorud, h.tarikh_khoruj, h.saat_khoruj])
                for h in items
            ),
        }
        valuation_overdue = bool(
            control.agreement_signed_on
            and not control.valuation_announced_on
            and (date.today() - control.agreement_signed_on).days > 30
        )
        try:
            saved_snapshot = (
                json.loads(control.final_summary_snapshot)
                if control.final_summary_snapshot else None
            )
        except (TypeError, json.JSONDecodeError):
            saved_snapshot = None
        return render_template(
            "reconciliation.html", control=control, items=items, stats=stats,
            valuation_overdue=valuation_overdue, saved_snapshot=saved_snapshot,
        )

    @reports_bp.route("/receipts/<path:filename>")
    @login_required
    def receipt_download(filename):
        if filename != secure_filename(filename) or "/" in filename or "\\" in filename:
            abort(404)
        return send_from_directory(
            UPLOAD_DIR, filename, as_attachment=True, max_age=0,
        )

    # ---------- error handlers ----------
    @app.errorhandler(400)
    def bad_request(error):
        return render_template(
            "error.html", code=400,
            message=getattr(error, "description", "درخواست معتبر نیست."),
        ), 400

    @app.errorhandler(403)
    def forbidden(_error):
        return render_template(
            "error.html", code=403,
            message="شما مجوز انجام این عملیات را ندارید.",
        ), 403

    @app.errorhandler(404)
    def not_found(_error):
        return render_template(
            "error.html", code=404,
            message="صفحه یا فایل موردنظر پیدا نشد.",
        ), 404

    @app.errorhandler(413)
    def too_large(_error):
        return render_template(
            "error.html", code=413,
            message="حجم فایل بیشتر از ۸ مگابایت است.",
        ), 413

    @app.errorhandler(503)
    def service_unavailable(error):
        return render_template(
            "error.html", code=503,
            message=getattr(
                error, "description",
                "سرویس موقتاً در دسترس نیست؛ دوباره تلاش کنید.",
            ),
        ), 503

    @app.errorhandler(500)
    def internal_error(_error):
        db.session.rollback()
        app.logger.exception("Internal error")
        return render_template(
            "error.html", code=500,
            message="خطای داخلی سرویس؛ موضوع به مدیر سامانه اطلاع داده شد.",
        ), 500

    # ---------- register ----------
    app.register_blueprint(auth_bp)
    app.register_blueprint(dashboard_bp)
    app.register_blueprint(havaleh_bp)
    app.register_blueprint(users_bp)
    app.register_blueprint(inventory_bp)
    app.register_blueprint(reports_bp)

    # ---------- startup bootstrap ----------
    with app.app_context():
        try:
            db.create_all()
        except OperationalError:
            app.logger.exception(
                "Database not available at startup; will retry on first request."
            )

        try:
            existing_columns = {
                c["name"] for c in sqlalchemy_inspect(db.engine).get_columns("havaleh")
            }
            for column_name in ("noe_vasileh", "pelak_trailer", "mogheiyat_varagh"):
                if column_name not in existing_columns:
                    column_type = (
                        "VARCHAR(200)" if column_name == "mogheiyat_varagh"
                        else "VARCHAR(30)"
                    )
                    db.session.execute(
                        text(f"ALTER TABLE havaleh ADD COLUMN {column_name} {column_type}")
                    )
            db.session.commit()
        except Exception:
            db.session.rollback()
            app.logger.exception("Schema check failed; continuing.")

        try:
            seed_path = BASE_DIR / "data" / "initial_inventory.csv"
            if seed_path.is_file() and not app.config.get("TESTING"):
                with seed_path.open("r", encoding="utf-8-sig", newline="") as seed_file:
                    for record in csv.DictReader(seed_file):
                        code = (record.get("code_kala") or "").strip()
                        spec = (record.get("moshakhasat") or "").strip() or None
                        location = (record.get("mogheiyat") or "").strip() or None
                        if not code or Inventory.query.filter_by(
                            code_kala=code, moshakhasat=spec, mogheiyat=location,
                        ).first():
                            continue

                        def seed_num(key):
                            raw = (record.get(key) or "").strip().replace(",", "")
                            return float(raw) if raw else None

                        db.session.add(Inventory(
                            radif=int(seed_num("radif") or 0) or None,
                            code_kala=code[:30],
                            name_kala=(record.get("name_kala") or "").strip()[:100] or None,
                            zekhamat=seed_num("zekhamat"),
                            moshakhasat=spec,
                            vahed=(record.get("vahed") or "کیلو گرم").strip()[:20],
                            mojudi_vaghei=seed_num("mojudi_vaghei"),
                            tedad_sheet=seed_num("tedad_sheet"),
                            vazn_har_sheet=seed_num("vazn_har_sheet"),
                            vazn_goruh=seed_num("vazn_goruh"),
                            mogheiyat=location,
                        ))
                db.session.commit()
        except Exception:
            db.session.rollback()
            app.logger.exception(
                "Initial inventory import failed; application startup will continue."
            )

        try:
            admin_phone = (os.environ.get("VARAGH_ADMIN_PHONE") or "").strip()
            bootstrap_password = os.environ.get("VARAGH_ADMIN_PASSWORD") or ""
            valid_admin_phone = bool(re.fullmatch(r"09\d{9}", admin_phone))

            if valid_admin_phone:
                phone_user = User.query.filter_by(username=admin_phone).first()
                if phone_user is None:
                    legacy_admin_name = os.environ.get("VARAGH_ADMIN_USERNAME", "admin")
                    legacy_admin = User.query.filter_by(
                        username=legacy_admin_name
                    ).first()
                    if legacy_admin is not None and legacy_admin.role == ROLE_ADMIN:
                        legacy_admin.username = admin_phone
                        if len(bootstrap_password) >= 12:
                            legacy_admin.set_password(bootstrap_password)
                        db.session.commit()

            if (
                User.query.count() == 0
                and bootstrap_password
                and len(bootstrap_password) >= 12
            ):
                initial_username = (
                    admin_phone if valid_admin_phone
                    else os.environ.get("VARAGH_ADMIN_USERNAME", "admin")
                )
                admin = User(
                    username=initial_username, full_name="مدیر سامانه",
                    role=ROLE_ADMIN, is_active_user=True,
                )
                admin.set_password(bootstrap_password)
                db.session.add(admin)
                db.session.commit()
        except Exception:
            db.session.rollback()
            app.logger.exception("Admin bootstrap failed; continuing startup.")

    return app


# ---------- Jalali conversions (module level) ----------
def gregorian_to_jalali(gy, gm, gd):
    """Convert Gregorian date to Jalali without an external dependency."""
    g_days_in_month = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    gy2 = gy + 1 if gm > 2 else gy
    days = (
        355666 + 365 * gy
        + (gy2 + 3) // 4 - (gy2 + 99) // 100 + (gy2 + 399) // 400
        + gd
    )
    for i in range(gm - 1):
        days += g_days_in_month[i]
    if gm > 2 and ((gy % 4 == 0 and gy % 100 != 0) or gy % 400 == 0):
        days += 1
    jy = -1595 + 33 * (days // 12053)
    days %= 12053
    jy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        jy += (days - 1) // 365
        days = (days - 1) % 365
    if days < 186:
        jm = 1 + days // 31
        jd = 1 + days % 31
    else:
        jm = 7 + (days - 186) // 30
        jd = 1 + (days - 186) % 30
    return jy, jm, jd


def jalali_to_gregorian(jy, jm, jd):
    """Convert a Jalali calendar date to Gregorian without third-party dependencies."""
    jy -= 979
    jm -= 1
    jd -= 1
    j_day_no = 365 * jy + (jy // 33) * 8 + ((jy % 33 + 3) // 4)
    for month in range(jm):
        j_day_no += 31 if month < 6 else 30
    j_day_no += jd
    g_day_no = j_day_no + 79
    gy = 1600 + 400 * (g_day_no // 146097)
    g_day_no %= 146097
    leap = True
    if g_day_no >= 36525:
        g_day_no -= 1
        gy += 100 * (g_day_no // 36524)
        g_day_no %= 36524
        if g_day_no >= 365:
            g_day_no += 1
        else:
            leap = False
    gy += 4 * (g_day_no // 1461)
    g_day_no %= 1461
    if g_day_no >= 366:
        leap = False
        g_day_no -= 1
        gy += g_day_no // 365
        g_day_no %= 365
    month_days = [31, 29 if leap else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    gm = 1
    while gm <= 12 and g_day_no >= month_days[gm - 1]:
        g_day_no -= month_days[gm - 1]
        gm += 1
    return date(gy, gm, g_day_no + 1)


def parse_flexible_date(raw):
    """Accept Jalali YYYY/MM/DD (Persian or Latin digits) and legacy ISO Gregorian dates."""
    raw = (raw or "").strip().translate(
        str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
    )
    if not raw:
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
        return date.fromisoformat(raw)
    match = re.fullmatch(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})", raw)
    if not match:
        raise ValueError(
            "تاریخ را به شکل شمسی سال/ماه/روز وارد کنید؛ نمونه: ۱۴۰۵/۰۷/۲۶"
        )
    jy, jm, jd = map(int, match.groups())
    if not (1 <= jm <= 12 and 1 <= jd <= (31 if jm <= 6 else 30 if jm <= 11 else 30)):
        raise ValueError("تاریخ شمسی واردشده معتبر نیست.")
    converted = jalali_to_gregorian(jy, jm, jd)
    if gregorian_to_jalali(converted.year, converted.month, converted.day) != (jy, jm, jd):
        raise ValueError("تاریخ شمسی واردشده معتبر نیست.")
    return converted


app = create_app()

if __name__ == "__main__":
    app.run(
        host=os.environ.get("HOST", "127.0.0.1"),
        port=int(os.environ.get("PORT", "5000")),
        debug=False,
    )
